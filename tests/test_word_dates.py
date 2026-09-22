import tempfile
import unittest
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from motore import word, word_inventory, word_patch
from motore.word_date import converti


class DatesTests(unittest.TestCase):
    def test_numeric_dates(self):
        self.assertEqual(converti('06/03/1971','dd-MM-yyyy'),('06-03-1971','1971-03-06T00:00:00Z'))
        self.assertEqual(converti('2024-02-29','d/M/yyyy')[0],'29/2/2024')
        for value in ('31/02/2024','02/03/24','ieri','', '2023-02-29'):
            with self.assertRaises(ValueError): converti(value)
        with self.assertRaises(ValueError): converti('2000-01-02','MMMM d, yyyy')

    def test_dates_and_existing_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); d=Document()
            for text, date, placeholder in [('Scegli data',True,True),('dfs',False,False),('Mario Rossi',False,False)]:
                p=d.add_paragraph('Campo: ')
                s=OxmlElement('w:sdt'); pr=OxmlElement('w:sdtPr'); s.append(pr)
                lock=OxmlElement('w:lock'); lock.set(qn('w:val'),'sdtLocked');pr.append(lock)
                if date:
                    dt=OxmlElement('w:date'); fmt=OxmlElement('w:dateFormat');fmt.set(qn('w:val'),'dd-MM-yyyy');dt.append(fmt);pr.append(dt)
                if placeholder:pr.append(OxmlElement('w:showingPlcHdr'))
                body=OxmlElement('w:sdtContent');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=text;r.append(t);body.append(r);s.append(body);p._p.append(s)
            src=root/'source.docx';d.save(src);m=word.leggi(src)
            self.assertEqual([a['scrivibile'] for a in m['ancore']],[True,False,False])
            vals={a['id']: '06/03/1971' for a in m['ancore']}
            out=root/'filled.docx';report=word.scrivi(src,m,vals,out)
            self.assertEqual(len(report['scritture']),1)
            self.assertEqual(len(report['mancate']),2)
            self.assertEqual(len(word.verifica(out,report)['verificate']),1)
            content,roots=word_inventory.apri(out.read_bytes())
            dt=roots['word/document.xml'].find('.//w:date',word_inventory.NS)
            self.assertEqual(dt.get(qn('w:fullDate')),'1971-03-06T00:00:00Z')
            texts=[b['testo'] for b in word_inventory.estrai(out.read_bytes())['blocchi']]
            self.assertEqual(texts,['Campo: 06-03-1971','Campo: dfs','Campo: Mario Rossi'])
            dt.set(qn('w:fullDate'),'2000-01-01T00:00:00Z')
            tampered=word_patch.salva(out.read_bytes(),{'word/document.xml':word_patch.serializza(roots['word/document.xml'])})
            self.assertFalse(word_patch.verifica_collocazione(src.read_bytes(),tampered,report['piano_audit'])['ok'])
            # Invalid input must leave the date control untouched.
            bad=root/'bad.docx'
            rejected=word.scrivi(src,m,{m['ancore'][0]['id']:'31/02/2024'},bad)
            self.assertFalse(rejected['scritture'])
            self.assertEqual(bad.read_bytes(),src.read_bytes())

    def test_block_date_control(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); d=Document(); p=d.add_paragraph('Scegli data')
            s=OxmlElement('w:sdt');pr=OxmlElement('w:sdtPr');s.append(pr)
            date=OxmlElement('w:date');fmt=OxmlElement('w:dateFormat')
            fmt.set(qn('w:val'),'yyyy-MM-dd');date.append(fmt);pr.append(date)
            pr.append(OxmlElement('w:showingPlcHdr'))
            content=OxmlElement('w:sdtContent');s.append(content)
            parent=p._p.getparent();index=list(parent).index(p._p)
            content.append(p._p);parent.insert(index,s)
            src=root/'source.docx';d.save(src);m=word.leggi(src)
            out=root/'out.docx';report=word.scrivi(src,m,{m['ancore'][0]['id']:'06/03/1971'},out)
            self.assertEqual(len(word.verifica(out,report)['verificate']),1)
