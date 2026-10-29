"""Il percorso Word "esperto": valori composti con fonti, righe senza spazio
segnato, controlli di provenienza e formato, revisione, memoria. Nessuna rete."""
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from motore import agente, compilatore_esperto as ce, lavorazione_word, word

PROFILO = {"ragione_sociale": "Aurora Servizi S.r.l.", "partita_iva": "01234567890",
           "nome_rappresentante": "Carla Ferrante", "luogo_nascita": "Imperia",
           "data_nascita": "14/05/1972", "codice_fiscale_rappresentante": "FRRCRL72E54E290X",
           "conto_dedicato_banca": "BANCA DI PROVA", "conto_dedicato_iban": "IT60 X054 2811 1010 0000 0123 456",
           "comune_sede_legale": "Imperia", "provincia_sede_legale": "IM"}
ENV = {"QUADRA_WORD_MODO": "esperto", "QUADRA_WORD_BASE_URL": "https://x.invalid/v1",
       "QUADRA_WORD_CHIAVE": "k", "QUADRA_WORD_MODELLO": "m"}


def modulo(d):
    d.add_paragraph("Il sottoscritto ______________ nato a ______________ il __________")
    d.add_paragraph("dell'impresa ____________________ P. IVA ____________")
    d.add_paragraph("gli estremi del conto sono:")
    d.add_paragraph("Banca")
    d.add_paragraph("IBAN")


class EspertoTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        doc = Document(); modulo(doc)
        self.src = self.root / "m.docx"; doc.save(self.src)
        self.rete = patch.object(agente, "_chiedi", side_effect=AssertionError("nessuna rete"))
        self.rete.start()

    def tearDown(self):
        self.rete.stop()
        self.temp.cleanup()

    def corri(self, risposte, scelte_di_prima=None, togli=()):
        domande = []
        risposte = list(risposte)

        def finto(testo, config, tetto=16000):
            domande.append(testo)
            return risposte.pop(0) if risposte else {"correzioni": [], "aggiunte": []}, 0.01
        with patch.dict(os.environ, ENV), patch.object(ce, "_chiedi", side_effect=finto):
            passi = list(lavorazione_word.lavora(self.root, self.src, {"profilo": PROFILO},
                                                 self.root / "esiti" / "chat_word",
                                                 pertinenza_confermata=True, togli=togli,
                                                 scelte_di_prima=scelte_di_prima))
        return passi[-1]["verdetto"], domande

    def numero(self, domanda, testo_riga, indice=0):
        riga = next(l for l in domanda.split("\n") if testo_riga in l)
        return riga.split("«")[1 + indice].split("»")[0].split(":")[0]

    def test_valori_composti_e_righe_senza_spazio(self):
        doc = Document(); modulo(doc); doc.save(self.root / "tmp.docx")
        m = word.leggi(self.root / "tmp.docx")
        testo, _ = ce.rendi(m)
        n = lambda r, i=0: self.numero(testo, r, i)
        banca = next(l for l in testo.split("\n") if l.endswith(" Banca")).split()[0][1:]
        proposta = {"campi": [
            {"posto": n("sottoscritto"), "valore": "Carla Ferrante", "fonti": ["nome_rappresentante"]},
            {"posto": n("sottoscritto", 1), "valore": "Imperia", "fonti": ["luogo_nascita"]},
            {"posto": n("sottoscritto", 2), "valore": "14/05/1972", "fonti": ["data_nascita"]},
            {"posto": n("impresa"), "valore": "Aurora Servizi S.r.l.", "fonti": ["ragione_sociale"]},
            {"posto": n("impresa", 1), "valore": "Genova", "fonti": ["luogo_nascita"]}],
            "aggiunte": [{"paragrafo": int(banca), "valore": "BANCA DI PROVA", "fonti": ["conto_dedicato_banca"]}]}
        v, domande = self.corri([proposta])
        scritti = {s["valore"] for s in v["verificate"]}
        self.assertEqual(scritti, {"Carla Ferrante", "Imperia", "14/05/1972", "Aurora Servizi S.r.l.", "BANCA DI PROVA"})
        self.assertIn("nome_rappresentante: Carla Ferrante", domande[0])
        scarti = {s["valore"]: s["motivo"] for s in v["esperto"]["scarti"]}
        self.assertIn("non viene dall'anagrafica", scarti["Genova"])
        self.assertEqual(len(domande), 2)  # proposta + un giro di revisione senza correzioni

    def test_formato_contro_il_testo_stampato(self):
        doc = Document(); modulo(doc); doc.save(self.root / "tmp.docx")
        testo, _ = ce.rendi(word.leggi(self.root / "tmp.docx"))
        piva = self.numero(testo, "impresa", 1)
        v, _ = self.corri([{"campi": [{"posto": piva, "valore": "Imperia", "fonti": ["comune_sede_legale"]}]}])
        self.assertFalse(v["verificate"])
        self.assertIn("partita IVA", v["esperto"]["scarti"][0]["motivo"])

    def test_la_revisione_corregge_e_poi_si_ferma(self):
        doc = Document(); modulo(doc); doc.save(self.root / "tmp.docx")
        testo, _ = ce.rendi(word.leggi(self.root / "tmp.docx"))
        nome = self.numero(testo, "sottoscritto")
        luogo = self.numero(testo, "sottoscritto", 1)
        v, domande = self.corri([
            {"campi": [{"posto": nome, "valore": "Imperia", "fonti": ["luogo_nascita"]}]},
            {"correzioni": [{"posto": nome, "valore": "Carla Ferrante", "fonti": ["nome_rappresentante"]},
                            {"posto": luogo, "valore": "Imperia", "fonti": ["luogo_nascita"]}]},
            {"correzioni": []}])
        self.assertEqual(sorted(s["valore"] for s in v["verificate"]), ["Carla Ferrante", "Imperia"])
        self.assertIn("⟦%s: Imperia⟧" % nome, domande[1])
        self.assertEqual(len(domande), 3)
        self.assertEqual(v["esperto"]["giri"], 2)

    def test_togliere_un_campo_non_richiama_il_modello_e_si_ricorda(self):
        doc = Document(); modulo(doc); doc.save(self.root / "tmp.docx")
        m = word.leggi(self.root / "tmp.docx")
        testo, ind = ce.rendi(m)
        banca = int(next(l for l in testo.split("\n") if l.endswith(" Banca")).split()[0][1:])
        nome = self.numero(testo, "sottoscritto")
        v, _ = self.corri([{"campi": [{"posto": nome, "valore": "Carla Ferrante", "fonti": ["nome_rappresentante"]}],
                            "aggiunte": [{"paragrafo": banca, "valore": "BANCA DI PROVA", "fonti": ["conto_dedicato_banca"]}]}])
        scelte = json.loads(json.dumps(v["scelte"]))
        aggiunta = next(a for a in scelte if a.endswith("_agg01"))
        v2, domande = self.corri([], scelte_di_prima=scelte, togli=[ind[nome]])
        self.assertEqual(domande, [])
        self.assertEqual([s["valore"] for s in v2["verificate"]], ["BANCA DI PROVA"])
        memoria = list((self.root / "esiti" / "memoria_moduli").glob("*.json"))
        self.assertEqual(len(memoria), 1)
        dati = json.loads(memoria[0].read_text(encoding="utf-8"))
        self.assertIn(ind[nome], dati["tolti"])
        self.assertIn(aggiunta, dati["confermati"])
        _, domande3 = self.corri([{"campi": []}])
        self.assertIn("chi compila l'ha tolto", domande3[0])

    def test_servizio_assente_torna_al_percorso_di_sempre(self):
        import requests
        with patch.dict(os.environ, ENV), \
                patch.object(ce, "_chiedi", side_effect=requests.ConnectionError("giu'")), \
                patch.object(agente, "abbina_pagina", return_value={}) as classico:
            passi = list(lavorazione_word.lavora(self.root, self.src, {"profilo": PROFILO},
                                                 self.root / "esiti" / "chat_word", pertinenza_confermata=True))
        self.assertTrue(classico.called)
        self.assertIn("percorso di sempre", " ".join(p["testo"] for p in passi))

    def test_controlli(self):
        self.assertIsNone(ce.controlla_provenienza(
            "nata a Imperia il 14/05/1972, c.f. FRRCRL72E54E290X",
            ["luogo_nascita", "data_nascita", "codice_fiscale_rappresentante"], PROFILO))
        self.assertIsNotNone(ce.controlla_provenienza("Imperia", [], PROFILO))
        self.assertIsNotNone(ce.controlla_provenienza("Imperia", ["inesistente"], PROFILO))
        self.assertIsNone(ce.controlla_formato("IBAN", "IT60 X054 2811 1010 0000 0123 456"))
        self.assertIsNotNone(ce.controlla_formato("IBAN", "IT61 X054 2811 1010 0000 0123 456"))
        self.assertIsNotNone(ce.controlla_formato("C.A.P.", "Imperia"))
        self.assertIsNone(ce.controlla_formato("con sede in Imperia (", "IM"))
        self.assertIsNotNone(ce.controlla_formato("PEC", "aurora.it"))
        self.assertIsNotNone(ce.sezione_esclusa(
            "[per gli operatori economici transfrontalieri] INDICA il seguente domicilio fiscale"))
        self.assertIsNone(ce.sezione_esclusa("con sede legale in"))


if __name__ == "__main__":
    unittest.main()
