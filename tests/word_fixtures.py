"""Synthetic four-field forms, created before any customer samples were supplied."""
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.enum.text import WD_TAB_LEADER
FIELDS = [("Ragione sociale", "Aurora Servizi S.r.l."), ("Partita IVA", "01234567890"), ("PEC", "aurora@example.invalid"), ("Sede", "Via delle Rose 12, Roma")]

def paragraphs(d, blank):
    for label, _ in FIELDS:
        d.add_paragraph(label + ': ' + blank)

def underline(d):
    for label, _ in FIELDS:
        p = d.add_paragraph(label + ': ')
        p.add_run(' ' * 25).underline = True

def tabs(d):
    for label, _ in FIELDS:
        p = d.add_paragraph(label + ':\t')
        p.paragraph_format.tab_stops.add_tab_stop(Inches(5), leader=WD_TAB_LEADER.DOTS)

def controls(d):
    for label, _ in FIELDS:
        p = d.add_paragraph(label + ': ')
        sdt = OxmlElement('w:sdt')
        props = OxmlElement('w:sdtPr')
        tag = OxmlElement('w:tag'); tag.set(qn('w:val'), label)
        props.append(tag); props.append(OxmlElement('w:text')); sdt.append(props)
        content = OxmlElement('w:sdtContent')
        r = OxmlElement('w:r'); t = OxmlElement('w:t')
        t.text = 'Fare clic qui per immettere testo.'
        r.append(t); content.append(r); sdt.append(content); p._p.append(sdt)

def table(d):
    t = d.add_table(rows=4, cols=2)
    for row, (label, _) in zip(t.rows, FIELDS):
        row.cells[0].text = label

def same_cell(d):
    t = d.add_table(rows=4, cols=1)
    for row, (label, _) in zip(t.rows, FIELDS):
        row.cells[0].text = label + ':\n'

def nested(d):
    outer = d.add_table(rows=1, cols=1)
    t = outer.cell(0,0).add_table(rows=4, cols=2)
    for row, (label, _) in zip(t.rows, FIELDS):
        row.cells[0].text = label

def header(d):
    for label, _ in FIELDS:
        d.sections[0].header.add_paragraph(label + ': _________________________')

def split_runs(d):
    for label, _ in FIELDS:
        p = d.add_paragraph(label + ': ')
        for v in ['____', '____', '____']:
            p.add_run(v).bold = True

def legacy(d):
    for label, _ in FIELDS:
        p = d.add_paragraph(label + ': ')
        start = OxmlElement('w:fldChar'); start.set(qn('w:fldCharType'),'begin')
        data = OxmlElement('w:ffData'); data.append(OxmlElement('w:textInput')); start.append(data)
        p.add_run()._r.append(start)
        instr = OxmlElement('w:instrText'); instr.text = ' FORMTEXT '
        p.add_run()._r.append(instr)
        sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'),'separate')
        p.add_run()._r.append(sep)
        p.add_run('\u2002'*5)
        end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end')
        p.add_run()._r.append(end)
