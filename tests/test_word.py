"""Offline regressions for generic Word structures and independent placement checks."""
import importlib.util
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from motore import word, word_inventory, word_patch, lavorazione_word, agente, soggetti
from word_fixtures import FIELDS, paragraphs, underline, tabs, controls, table, same_cell, nested, header, split_runs, legacy


class WordTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        # Any unmocked model call is a test error, never an API purchase.
        self.network = patch.object(agente, '_chiedi', side_effect=AssertionError('No network in tests'))
        self.network.start()

    def tearDown(self):
        self.network.stop()
        self.temp.cleanup()

    def source(self, build, name='source'):
        doc = Document(); build(doc)
        path = self.root / (name+'.docx'); doc.save(path)
        return path

    def fill(self, source, values=None):
        mapping=word.leggi(source)
        values=values if values is not None else {a['id']:FIELDS[i][1] for i,a in enumerate(mapping['ancore'])}
        output=self.root/'filled.docx'
        report=word.scrivi(source,mapping,values,output)
        return mapping,output,report,word.verifica(output,report)

    def test_four_fields_across_fourteen_structures(self):
        cases = dict(underscores=lambda d:paragraphs(d,'_'*25),
                     dots=lambda d:paragraphs(d,'.'*25), hyphens=lambda d:paragraphs(d,'-'*25),
                     underlined_spaces=underline, tab_leaders=tabs, sdt=controls,
                     empty_cells=table, same_cell=same_cell, nested=nested, header=header,
                     split_runs=split_runs, formtext=legacy,
                     short_underscores=lambda d:paragraphs(d,'___'),
                     double_ellipsis=lambda d:paragraphs(d,'……'))
        for name,build in cases.items():
            with self.subTest(name=name):
                src=self.source(build,name)
                mapping,out,report,check=self.fill(src)
                self.assertEqual(len(mapping['ancore']),4)
                self.assertEqual(len(report['scritture']),4)
                self.assertEqual(len(check['verificate']),4)
                self.assertFalse(check['non_trovate'])
                self.assertFalse(check['mancate'])
                text='\n'.join(b['testo'] for b in word_inventory.estrai(out.read_bytes())['blocchi'])
                for _,value in FIELDS:
                    self.assertIn(value,text)
                self.assertEqual(check['verifica_visiva'],'NON_ESEGUITA')

    def test_repeated_address_stays_complete(self):
        value='Via delle Rose 12 Roma'
        src=self.source(lambda d:d.add_paragraph('Residenza: '+value+'; Sede legale: ___'))
        m=word.leggi(src)
        _,out,_,check=self.fill(src,{m['ancore'][0]['id']:value})
        self.assertEqual(Document(out).paragraphs[0].text.count(value),2)
        self.assertEqual(len(check['verificate']),1)

    def test_single_cell_labels_without_placeholders(self):
        labels = ['Nome e Cognome', 'Codice fiscale', 'Ragione sociale.', 'Identificativo pratica']
        def build(d):
            t = d.add_table(rows=4, cols=1)
            for row, label in zip(t.rows, labels):
                row.cells[0].text = label
        src = self.source(build)
        m = word.leggi(src)
        self.assertEqual(len(m['ancore']), 4)
        self.assertTrue(all(a['etichetta_implicita'] for a in m['ancore']))
        text, addresses = word.rendi(m)
        self.assertEqual(len(addresses), 4)
        self.assertIn('inserimento facoltativo', text)
        values = {a['id']: 'Dato sintetico %d' % i for i, a in enumerate(m['ancore'])}
        _, out, _, check = self.fill(src, values)
        self.assertEqual(len(check['verificate']), 4)
        for i, row in enumerate(Document(out).tables[0].rows):
            self.assertEqual(row.cells[0].text, labels[i] + ' Dato sintetico %d' % i)

    def test_table_title_candidate_is_never_automatically_filled(self):
        src = self.source(lambda d: setattr(d.add_table(rows=1, cols=1).cell(0,0), 'text', 'DICHIARAZIONI'))
        m, out, report, check = self.fill(src, {})
        self.assertEqual(len(m['ancore']), 1)
        self.assertFalse(report['scritture'])
        self.assertEqual(Document(out).tables[0].cell(0,0).text, 'DICHIARAZIONI')

    def test_arbitrary_label_cannot_truncate_company_name(self):
        src=self.source(lambda d:d.add_paragraph('Società: ___'))
        m=word.leggi(src)
        _,out,_,_=self.fill(src,{m['ancore'][0]['id']:'Società Agricola Aurora S.r.l.'})
        self.assertIn('Società Agricola Aurora S.r.l.',Document(out).paragraphs[0].text)

    def test_order_prevents_wrong_section_exclusion(self):
        def build(d):
            d.add_paragraph('DATI IMPRESA CONCORRENTE'); table(d)
            d.add_paragraph('DATI IMPRESA AUSILIARIA - AVVALIMENTO')
        src=self.source(build); m=word.leggi(src); text,addresses=word.rendi(m)
        self.assertLess(text.index('«1:'),text.index('DATI IMPRESA AUSILIARIA'))
        self.assertFalse(soggetti.da_saltare(lavorazione_word._finta_mappa(text,addresses),set()))

    def test_swapped_values_fail_position_verification(self):
        src=self.source(lambda d:d.add_paragraph('Nome: ___ Cognome: ___'))
        m=word.leggi(src)
        _,out,report,check=self.fill(src,dict(zip([a['id'] for a in m['ancore']],['Mario','Rossi'])))
        self.assertEqual(len(check['verificate']),2)
        changed=Document(out); changed.paragraphs[0].text='Nome: Rossi Cognome: Mario'; changed.save(out)
        check=word.verifica(out,report)
        self.assertFalse(check['verificate']); self.assertEqual(len(check['non_trovate']),2)

    def test_unauthorized_header_change_fails_verification(self):
        src=self.source(lambda d:d.add_paragraph('Nome: ___'))
        m=word.leggi(src); _,out,report,_=self.fill(src,{m['ancore'][0]['id']:'Mario'})
        changed=Document(out); changed.sections[0].header.paragraphs[0].text='Alterato'; changed.save(out)
        self.assertFalse(word.verifica(out,report)['verificate'])

    def test_changed_source_rejected(self):
        src=self.source(lambda d:paragraphs(d,'___')); m=word.leggi(src)
        d=Document(src); d.add_paragraph('Nuovo testo'); d.save(src)
        with self.assertRaisesRegex(ValueError,'cambiato'):
            word.scrivi(src,m,{},self.root/'filled.docx')

    def test_label_without_drawn_placeholder(self):
        src=self.source(lambda d:[d.add_paragraph(label+':') for label,_ in FIELDS])
        _,_,_,check=self.fill(src)
        self.assertEqual(len(check['verificate']),4)

    def test_locked_sdt_is_reported_and_not_written(self):
        def build(d):
            controls(d)
            for control in d.element.iter(qn('w:sdt')):
                lock=OxmlElement('w:lock'); lock.set(qn('w:val'),'contentLocked')
                control.find(qn('w:sdtPr')).append(lock)
        src=self.source(build); m=word.leggi(src)
        self.assertEqual(len(m['ancore']),4)
        self.assertTrue(all(not a['scrivibile'] for a in m['ancore']))
        self.assertFalse(word.rendi(m)[1])
        _,_,_,check=self.fill(src)
        self.assertFalse(check['verificate']); self.assertEqual(len(check['mancate']),4)

    def test_plain_rich_text_and_split_sdt_results(self):
        def build(d):
            controls(d)
            for control in d.element.iter(qn('w:sdt')):
                pr=control.find(qn('w:sdtPr')); pr.remove(pr.find(qn('w:text')))
                content=control.find(qn('w:sdtContent'))
                r=OxmlElement('w:r'); t=OxmlElement('w:t'); t.text=' Altro frammento'
                r.append(t); content.append(r)
        src=self.source(build)
        _,_,_,check=self.fill(src)
        self.assertEqual(len(check['verificate']),4)

    def test_empty_inline_sdt_can_be_filled(self):
        def build(d):
            controls(d)
            for content in d.element.iter(qn('w:sdtContent')):
                for node in list(content): content.remove(node)
        src=self.source(build)
        _,_,_,check=self.fill(src)
        self.assertEqual(len(check['verificate']),4)

    def test_footer_fields(self):
        def build(d):
            for label,_ in FIELDS:
                d.sections[0].footer.add_paragraph(label+': ___')
        _,_,_,check=self.fill(self.source(build))
        self.assertEqual(len(check['verificate']),4)

    def test_xml_textbox_is_not_counted_twice(self):
        def build(d):
            p=d.add_paragraph('Modulo')
            box=OxmlElement('w:txbxContent')
            for label,_ in FIELDS:
                para=OxmlElement('w:p'); r=OxmlElement('w:r'); t=OxmlElement('w:t')
                t.text=label+': ___'; r.append(t); para.append(r); box.append(para)
            pict=OxmlElement('w:pict'); pict.append(box); p.add_run()._r.append(pict)
        _,_,_,check=self.fill(self.source(build))
        self.assertEqual(len(check['verificate']),4)

    def test_cross_run_placeholder_preserves_bold_label(self):
        def build(d):
            p=d.add_paragraph(); p.add_run('Nome: ').bold=True
            p.add_run('__').italic=True; p.add_run('___')
        src=self.source(build); m=word.leggi(src)
        _,out,_,_=self.fill(src,{m['ancore'][0]['id']:'Mario'})
        p=Document(out).paragraphs[0]
        self.assertTrue(p.runs[0].bold); self.assertTrue(p.runs[1].italic)
        self.assertEqual(p.text,'Nome: Mario')

    def test_underlined_blank_keeps_its_format_and_label_spacing(self):
        src=self.source(underline)
        _,out,_,_=self.fill(src)
        p=Document(out).paragraphs[0]
        self.assertEqual(p.text,FIELDS[0][0]+': '+FIELDS[0][1])
        self.assertTrue(next(r for r in p.runs if FIELDS[0][1] in r.text).underline)

    def test_mixed_fields_no_false_one_out_of_one(self):
        def build(d):
            d.add_paragraph(FIELDS[0][0]+': ___')
            for label,_ in FIELDS[1:]:
                d.add_paragraph(label+': ').add_run(' '*20).underline=True
        src=self.source(build); m=word.leggi(src)
        self.assertEqual(len(m['ancore']),4)

    def test_workflow_mapping_review_export(self):
        src=self.source(controls); m=word.leggi(src)
        profile={'profilo':dict(zip(['ragione_sociale','partita_iva','pec','sede'],[v for _,v in FIELDS]))}
        decisions=dict(zip([a['id'] for a in m['ancore']],profile['profilo']))
        with patch.object(agente,'abbina_pagina',return_value=decisions), patch.object(agente,'rileggi_pagina',return_value={}):
            events=list(lavorazione_word.lavora(self.root,src,profile,self.root/'work',pertinenza_confermata=True))
        verdict=events[-1]['verdetto']
        self.assertEqual(len(verdict['verificate']),4)
        self.assertEqual(verdict['esito'],'compilato')
        self.assertEqual(verdict['assegnati'],4)
        self.assertTrue(Path(verdict['bozza']).exists())
        self.assertIn('valori', verdict['tempi_secondi'])
        self.assertIn('correzione', verdict['tempi_secondi'])
        self.assertIn('scrittura', verdict['tempi_secondi'])
        logs = list((self.root/'work').glob('diagnostica_word_*.json'))
        self.assertEqual(len(logs), 1)
        for value in profile['profilo'].values():
            self.assertNotIn(value, logs[0].read_text(encoding='utf-8'))

    def test_empty_mapping_and_review_failure_are_not_success(self):
        src=self.source(lambda d:paragraphs(d,'___')); m=word.leggi(src)
        profile={'profilo':dict(zip(['ragione_sociale','partita_iva','pec','sede'],[v for _,v in FIELDS]))}
        decisions=dict(zip([a['id'] for a in m['ancore']],profile['profilo']))
        for mapping,review_failure in [({},False),(decisions,True)]:
            with self.subTest(review_failure=review_failure):
                with patch.object(agente,'abbina_pagina',return_value=mapping), patch.object(agente,'rileggi_pagina',side_effect=RuntimeError('Offline fault')):
                    events=list(lavorazione_word.lavora(self.root,src,profile,self.root/'work',pertinenza_confermata=True))
                verdict=events[-1]['verdetto']
                self.assertEqual(verdict['esito'],'non_compilato')
                self.assertEqual(bool(verdict['guasti']),review_failure)

    def test_legacy_regressions(self):
        path=Path(__file__).with_name('legacy_word_checks.py')
        spec=importlib.util.spec_from_file_location('legacy_word_checks',path)
        module=importlib.util.module_from_spec(spec); spec.loader.exec_module(module)
        module.TMP=self.root/'legacy'
        for name,func in module._prove:
            with self.subTest(name=name):
                func()


if __name__=='__main__':
    unittest.main()
