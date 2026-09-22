# -*- coding: utf-8 -*-
"""Le prove del percorso Word.

Si costruiscono i documenti qui dentro invece di prenderli dal disco: i moduli
veri stanno in `corpus_word/`, che non entra nel repository, e una prova che
dipende da file assenti non gira per chi prende il progetto.

    python -m prove.prova_word
"""
from __future__ import annotations

import io

from motore import word

_prove = []


def prova(descrizione):
    def decoratore(f):
        _prove.append((descrizione, f))
        return f
    return decoratore


def _documento(righe, tabella=None):
    """Un .docx finto, in memoria, con i paragrafi dati."""
    import docx
    d = docx.Document()
    for riga in righe:
        d.add_paragraph(riga)
    if tabella:
        t = d.add_table(rows=len(tabella), cols=len(tabella[0]))
        for i, riga in enumerate(tabella):
            for j, cella in enumerate(riga):
                t.cell(i, j).text = cella
    fuori = io.BytesIO()
    d.save(fuori)
    return fuori.getvalue()


def _con_campi_modulo(pezzi):
    """Un .docx con dentro i campi modulo veri di Word (FORMTEXT).

    Sono la cosa che distingue un modulo fatto bene da uno fatto coi puntini,
    e si costruiscono a mano perche' python-docx non sa farli: begin con
    ffData/textInput, il codice, separate, il risultato - cinque spazi EN,
    quelli che ci mette Word - e end.

    `pezzi` alterna testo e None, dove None e' un campo.
    """
    import docx
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    d = docx.Document()
    p = d.add_paragraph()
    for pezzo in pezzi:
        if pezzo is not None:
            p.add_run(pezzo)
            continue
        for xml in (
            '<w:r %s><w:fldChar w:fldCharType="begin"><w:ffData>'
            '<w:name w:val="Testo"/><w:enabled/><w:textInput/>'
            '</w:ffData></w:fldChar></w:r>',
            '<w:r %s><w:instrText xml:space="preserve"> FORMTEXT </w:instrText></w:r>',
            '<w:r %s><w:fldChar w:fldCharType="separate"/></w:r>',
            '<w:r %s><w:t xml:space="preserve">     </w:t></w:r>',
            '<w:r %s><w:fldChar w:fldCharType="end"/></w:r>',
        ):
            p._p.append(parse_xml(xml % nsdecls("w")))
    fuori = io.BytesIO()
    d.save(fuori)
    return fuori.getvalue()


def _leggi(byte, tmp):
    percorso = tmp / "finto.docx"
    percorso.parent.mkdir(parents=True, exist_ok=True)
    percorso.write_bytes(byte)
    return word.leggi(percorso, tmp)


import pathlib
TMP = pathlib.Path("esiti") / "prova_word"


@prova("una riga di trattini diventa un campo con la sua capienza")
def _riempimenti():
    m = _leggi(_documento(["Il sottoscritto ______________ nato a _______ il ____"]), TMP)
    riemp = [a for a in m["ancore"] if a["tipo"] == "riempimento"]
    if len(riemp) != 3:
        raise AssertionError("attesi 3 campi, trovati %d" % len(riemp))
    if not (riemp[0]["capienza"] > riemp[1]["capienza"] > riemp[2]["capienza"]):
        raise AssertionError("le capienze non seguono la lunghezza dei trattini: %s"
                             % [a["capienza"] for a in riemp])


@prova("l'etichetta si taglia al campo precedente")
def _etichette():
    m = _leggi(_documento(["Il sottoscritto ______ nato a ______"]), TMP)
    etichette = [a["etichetta"] for a in m["ancore"]]
    if "nato a" not in etichette[1]:
        raise AssertionError("seconda etichetta sbagliata: %r" % etichette[1])
    if "sottoscritto" in etichette[1]:
        raise AssertionError("l'etichetta si e' allungata sul campo prima: %r"
                             % etichette[1])


@prova("puntini e puntini di sospensione valgono come trattini")
def _altri_riempitivi():
    m = _leggi(_documento(["Comune ........... Provincia ………"]), TMP)
    if len([a for a in m["ancore"] if a["tipo"] == "riempimento"]) != 2:
        raise AssertionError("puntini non riconosciuti: %s"
                             % [a["tipo"] for a in m["ancore"]])


@prova("i campi dentro le celle di tabella non si perdono")
def _celle():
    m = _leggi(_documento(["testo fuori"],
                          tabella=[["Ragione sociale", "______________"]]), TMP)
    dentro = [a for a in m["ancore"] if a["tabella"] is not None]
    if not dentro:
        raise AssertionError("nessun campo trovato nelle celle")


@prova("la resa usa gli stessi buchi numerati del percorso PDF")
def _resa():
    m = _leggi(_documento(["Il sottoscritto ______ nato a ______"]), TMP)
    testo, indirizzi = word.rendi(m)
    if "«1:" not in testo or "«2:" not in testo:
        raise AssertionError("mancano i buchi numerati: %r" % testo)
    if "Il sottoscritto" not in testo or "nato a" not in testo:
        raise AssertionError("la frase intorno si e' persa: %r" % testo)
    if set(indirizzi) != {"1", "2"}:
        raise AssertionError("indirizzi sbagliati: %s" % indirizzi)


@prova("un paragrafo senza campi resta nella prosa, per fare contesto")
def _contesto():
    m = _leggi(_documento(["L'IMPRESA", "Ragione sociale ______"]), TMP)
    testo, _ = word.rendi(m)
    if "L'IMPRESA" not in testo:
        raise AssertionError("l'intestazione senza campi e' sparita: %r" % testo)


@prova("una casella si riconosce e si numera come tale")
def _caselle():
    m = _leggi(_documento(["☐ Legale rappresentante", "☐ Procuratore"]), TMP)
    if len([a for a in m["ancore"] if a["tipo"] == "casella"]) != 2:
        raise AssertionError("caselle non riconosciute")
    testo, _ = word.rendi(m)
    if "«1:casella»" not in testo:
        raise AssertionError("la casella non e' segnata come tale: %r" % testo)
    if "Legale rappresentante" not in testo:
        raise AssertionError("l'etichetta della casella si e' persa")


@prova("celle vuote e trattini si scrivono insieme senza inciampare")
def _celle_e_trattini():
    byte = _documento(["Il sottoscritto ______ nato a ______"],
                      tabella=[["Ragione sociale", ""], ["Partita IVA", ""]])
    m = _leggi(byte, TMP)
    celle = [a for a in m["ancore"] if a["tipo"] == "cella"]
    if not celle:
        raise AssertionError("le celle vuote non sono state viste come campi")
    valori = {a["id"]: "X%d" % i for i, a in enumerate(m["ancore"])}
    # Il guasto vero: le celle hanno paragrafo None e ordinarle insieme agli
    # altri campi confronta None con un numero. Due moduli su sei si fermavano.
    r = word.scrivi(TMP / "finto.docx", m, valori, TMP / "bozza_mista.docx")
    v = word.verifica(TMP / "bozza_mista.docx", r)
    if v["non_trovate"]:
        raise AssertionError("scritture non ritrovate: %s" % v["non_trovate"])
    if len(v["verificate"]) != len(valori):
        raise AssertionError("attese %d scritture, verificate %d"
                             % (len(valori), len(v["verificate"])))


@prova("una riga di elenco si riconosce come roba di altri soggetti")
def _elenchi():
    m = _leggi(_documento(["testo"],
                          tabella=[["Ragione Sociale", "C.F."],
                                   ["", ""], ["", ""], ["", ""]]), TMP)
    elenco = [a for a in m["ancore"] if a.get("elenco")]
    if len(elenco) < 4:
        raise AssertionError("le righe sotto l'intestazione non sono segnate "
                             "come elenco: %s"
                             % [(a["id"], a.get("elenco")) for a in m["ancore"]])
    testo, _ = word.rendi(m)
    if "elenco di altri soggetti" not in testo:
        raise AssertionError("la resa non avverte il modello: %r" % testo[:200])


@prova("alla rilettura si passano le chiavi, non i valori")
def _rilettura_vuole_chiavi():
    """Il guasto: passandole i valori, la rilettura cercava
    `anagrafe.valore("COSTRUZIONI RIVAMARE S.R.L.")`, trovava None e bocciava
    ogni scrittura con "None non e' un valore valido". Su un modulo ha
    azzerato tutte e tre le scritture, che erano corrette, e nessun conteggio
    lo diceva: il verdetto riportava zero scritture e zero errori."""
    import inspect
    from motore import agente, lavorazione_word
    sorgente = inspect.getsource(lavorazione_word.lavora)
    dove = sorgente.find("rileggi_documento")
    if dove < 0:
        raise AssertionError("il percorso Word non rilegge piu' prima di scrivere")
    # la chiamata va a capo: si guarda un pezzo di testo, non la riga sola
    if "chiavi_scelte" not in sorgente[dove:dove + 160]:
        raise AssertionError("alla rilettura non si passano le chiavi:\n%s"
                             % sorgente[dove:dove + 160])
    # e la firma dall'altra parte non deve cambiare sotto i piedi
    parametri = list(inspect.signature(agente.rileggi_documento).parameters)
    if parametri[:3] != ["rese", "scelte", "anagrafe"]:
        raise AssertionError("rileggi_documento ha cambiato firma: %s" % parametri)


@prova("un vuoto disegnato con segni diversi resta un campo solo")
def _vuoto_misto():
    """Il guasto: "cap ………..... Via" e' un buco solo, scritto con tre puntini
    di sospensione e cinque punti. Le tre espressioni separate lo spezzavano
    in due campi attaccati, il modello ne vedeva due e li riempiva tutti e
    due: usciva "cap IM 18100", con la provincia dentro il C.A.P. Misurato
    sul corpus: 215 campi su 2460, in 18 moduli su 48."""
    m = _leggi(_documento(["cap ………..... Via ......……… n. ___"]), TMP)
    riemp = [a for a in m["ancore"] if a["tipo"] == "riempimento"]
    if len(riemp) != 3:
        raise AssertionError(
            "attesi 3 campi (cap, Via, n.), trovati %d: %s"
            % (len(riemp), [(a["etichetta"], a["capienza"]) for a in riemp]))
    if "cap" not in riemp[0]["etichetta"].lower():
        raise AssertionError("primo campo non e' il cap: %r" % riemp[0]["etichetta"])
    if riemp[0]["capienza"] != 8:
        raise AssertionError("la capienza non copre tutto il tratto: %d"
                             % riemp[0]["capienza"])


@prova("una punteggiatura normale non diventa un campo")
def _punteggiatura():
    m = _leggi(_documento(["Lavori vari... e altro. Importo ___"]), TMP)
    riemp = [a for a in m["ancore"] if a["tipo"] == "riempimento"]
    if len(riemp) != 1:
        raise AssertionError("i puntini di sospensione della prosa sono "
                             "diventati campi: %s"
                             % [(a["etichetta"], a["capienza"]) for a in riemp])


@prova("la parola gia' stampata sul modulo non si riscrive")
def _niente_via_via():
    """Sul foglio c'e' "Via ........." e l'indirizzo e' "Via delle Fornaci":
    scritto intero si legge "Via Via delle Fornaci". Misurato: 20 volte su
    654 scritture, tutte "Via"."""
    m = _leggi(_documento(["Via ______________ n. ____",
                           "Comune ______________"]), TMP)
    per_etichetta = {a["etichetta"]: a["id"] for a in m["ancore"]}
    valori = {per_etichetta["Via"]: "Via delle Fornaci",
              per_etichetta["n"]: "9/A",
              per_etichetta["Comune"]: "Imperia"}
    r = word.scrivi(TMP / "finto.docx", m, valori, TMP / "bozza_via.docx")
    import docx
    righe = [p.text for p in docx.Document(str(TMP / "bozza_via.docx")).paragraphs]
    if "Via Via" in "\n".join(righe):
        raise AssertionError("la parola si e' ripetuta: %r" % righe[0])
    if "Via delle Fornaci" not in "Via " + righe[0]:
        raise AssertionError("la strada si e' persa: %r" % righe[0])
    # e un valore che non comincia con la parola dell'etichetta resta intero
    if "Imperia" not in righe[1]:
        raise AssertionError("il comune si e' perso: %r" % righe[1])


@prova("una frase ripetuta non autorizza a troncare il dato sorgente")
def _niente_eco():
    """La castroneria piu' frequente rimasta dopo il collaudo: dieci
    scritture su 668, tutte questa.

        ...registro delle imprese della Camera di Commercio Industria
        Artigianato ed Agricoltura di [Camera di Commercio Riviere di
        Liguria - Imperia La Spezia Savona]
    """
    riga = ("che l'impresa e' iscritta nel registro delle imprese della "
            "Camera di Commercio Industria Artigianato ed Agricoltura di ")
    valore = "Camera di Commercio Riviere di Liguria - Imperia La Spezia Savona"
    fuori = word._senza_eco(riga, valore)
    if fuori != valore:
        raise AssertionError("il dato sorgente e' stato alterato: %r" % fuori)
    if "Riviere di Liguria" not in fuori:
        raise AssertionError("ha tagliato troppo: %r" % fuori)


@prova("l'eco non tocca quello che romperebbe")
def _eco_prudente():
    """Le tre cose su cui la versione larga di questa regola faceva danni, e
    che le due soglie - frase di almeno due parole, resto di almeno due -
    tengono fuori. Misurate: con soglie a uno scattava 25 volte e faceva
    disastri, con soglie a due scatta 9 volte e sono tutte giuste."""
    casi = [
        ("tel 0183 700111 e-mail rivamare@pec.example.invalid PEC ",
         "rivamare@pec.example.invalid", "una sola parola ripetuta"),
        ("in qualita' di legale rappresentante dell'impresa ",
         "Legale rappresentante", "non resterebbe niente"),
        ("SOCIETA' (specificare tipo) ",
         "Societa' a responsabilita' limitata", "una sola parola ripetuta"),
        ("in qualita' di (titolare, legale rappresentante, procuratore) ",
         "Legale rappresentante", "fra parentesi c'e' un elenco di scelte"),
    ]
    for riga, valore, perche in casi:
        if word._senza_eco(riga, valore) != valore:
            raise AssertionError("ha toccato %r (%s) -> %r"
                                 % (valore, perche, word._senza_eco(riga, valore)))


@prova("una parola corta uguale all'etichetta non si tocca")
def _parole_corte():
    if word._senza_ripetizione("n.", "n 9/A") != "n 9/A":
        raise AssertionError("ha tolto una parola di due lettere")
    if word._senza_ripetizione("Via", "Via") != "Via":
        raise AssertionError("ha svuotato un valore di una parola sola")
    if word._senza_ripetizione("Comune di", "Imperia") != "Imperia":
        raise AssertionError("ha toccato un valore che non ripete niente")


@prova("togliere un campo sbagliato non muove gli altri")
def _togliere_non_devasta():
    """La domanda vera di chi usera' l'app: se una castroneria si puo'
    togliere senza rovinare il resto. Rifare la bozza richiamava il modello, e
    la risposta e' libera di cambiare. Adesso si riusa la decisione gia'
    presa, e questa prova lo inchioda confrontando le due bozze riga per riga:
    ne deve cambiare una sola.

    Gira senza modello: `scelte_di_prima` e `pertinenza_confermata` tolgono
    tutte e due le chiamate, quindi si puo' ripetere a ogni modifica."""
    import docx
    from motore import lavorazione_word

    cartella = TMP / "togliere"
    percorso = cartella / "modulo.docx"
    percorso.parent.mkdir(parents=True, exist_ok=True)
    percorso.write_bytes(_documento([
        "Il sottoscritto ____________ nato a __________ il __________",
        "in qualita' di ____________ dell'impresa ____________________",
        "con sede in ____________ C.A.P. ______ tel. ____________"]))

    mappa = word.leggi(percorso, cartella / "convertiti")
    _, indirizzi = word.rendi(mappa)
    ancore = [indirizzi[str(n)] for n in range(1, len(indirizzi) + 1)]
    chiavi = ["nome_cognome", "comune_nascita", "data_nascita", "carica",
              "denominazione", "comune_sede", "cap_sede", "telefono"]
    fascicolo = {"profilo": {"nome_cognome": "Carla Ferrante",
                             "comune_nascita": "Imperia",
                             "data_nascita": "14/05/1972",
                             "carica": "Legale rappresentante",
                             "denominazione": "COSTRUZIONI RIVAMARE S.R.L.",
                             "comune_sede": "Imperia", "cap_sede": "18100",
                             "telefono": "0183 700111"}}
    scelte = dict(zip(ancore, chiavi))

    def bozza(togli, dove):
        ultimo = None
        for passo in lavorazione_word.lavora(
                TMP, percorso, fascicolo, cartella / dove,
                pertinenza_confermata=True, togli=togli, condizioni=set(),
                scelte_di_prima=scelte):
            if passo.get("verdetto"):
                ultimo = passo["verdetto"]
        if ultimo is None:
            raise AssertionError("nessun verdetto da %s" % dove)
        return ultimo

    intera = bozza((), "intera")
    if len(intera["verificate"]) != len(scelte):
        raise AssertionError("la bozza intera ha scritto %d campi su %d"
                             % (len(intera["verificate"]), len(scelte)))

    # si toglie quello in mezzo, che e' il caso peggiore: spostando il testo
    # cambiano le posizioni di tutti i campi che seguono nello stesso paragrafo
    tolta = ancore[3]
    ridotta = bozza([tolta], "ridotta")

    rimaste = {r["ancora"]: r["valore"] for r in ridotta["verificate"]}
    prima = {r["ancora"]: r["valore"] for r in intera["verificate"]}
    if tolta in rimaste:
        raise AssertionError("il campo spuntato e' stato scritto lo stesso")
    if set(prima) - {tolta} != set(rimaste):
        raise AssertionError(
            "togliendo un campo se ne sono mossi altri: spariti %s, comparsi %s"
            % (sorted(set(prima) - {tolta} - set(rimaste)),
               sorted(set(rimaste) - set(prima))))
    diversi = {a: (prima[a], rimaste[a]) for a in rimaste if prima[a] != rimaste[a]}
    if diversi:
        raise AssertionError("valori cambiati negli altri campi: %s" % diversi)

    # e il documento vero: identico, tranne il valore tolto
    def righe(v):
        return [p.text for p in docx.Document(v["bozza"]).paragraphs]
    ri, rr = righe(intera), righe(ridotta)
    if len(ri) != len(rr):
        raise AssertionError("le due bozze non hanno lo stesso numero di righe")
    cambiate = [(a, b) for a, b in zip(ri, rr) if a != b]
    if len(cambiate) != 1:
        raise AssertionError("attesa una riga cambiata, cambiate %d: %s"
                             % (len(cambiate), cambiate))
    if prima[tolta] in cambiate[0][1]:
        raise AssertionError("il valore tolto e' ancora nella riga: %r"
                             % cambiate[0][1])


@prova("i campi modulo veri di Word si vedono e si riempiono")
def _campi_modulo():
    """Un modulo fatto bene non usa i puntini: usa i campi FORMTEXT. Non li
    vedevamo, e i moduli fatti bene erano quelli che ci riuscivano peggio.
    Misurato sul corpus: 11 documenti su 48 li usano, e in tre erano l'unico
    modo di riempirli - 261, 99 e 72 campi. Quei tre scrivevano zero."""
    byte = _con_campi_modulo(["Il sottoscritto ", None, " nato a ", None])
    percorso = TMP / "modulo_word.docx"
    percorso.parent.mkdir(parents=True, exist_ok=True)
    percorso.write_bytes(byte)
    m = word.leggi(percorso, TMP)
    campi = [a for a in m["ancore"] if a["tipo"] == "modulo"]
    if len(campi) != 2:
        raise AssertionError("attesi 2 campi modulo, trovati %d su %s"
                             % (len(campi), [a["tipo"] for a in m["ancore"]]))
    if "sottoscritto" not in campi[0]["etichetta"]:
        raise AssertionError("prima etichetta sbagliata: %r" % campi[0]["etichetta"])
    if campi[1]["etichetta"].strip() != "nato a":
        raise AssertionError("l'etichetta si e' allungata sul campo prima: %r"
                             % campi[1]["etichetta"])
    testo, indirizzi = word.rendi(m)
    if "«1:" not in testo or "«2:" not in testo:
        raise AssertionError("i campi modulo non entrano nella prosa: %r" % testo)

    # e ci si scrive dentro con la stessa strada dei puntini
    valori = {campi[0]["id"]: "Carla Ferrante", campi[1]["id"]: "Imperia"}
    r = word.scrivi(percorso, m, valori, TMP / "bozza_modulo_word.docx")
    v = word.verifica(TMP / "bozza_modulo_word.docx", r)
    if v["non_trovate"] or len(v["verificate"]) != 2:
        raise AssertionError("scritture non ritrovate: %s" % v["non_trovate"])
    import docx
    riga = docx.Document(str(TMP / "bozza_modulo_word.docx")).paragraphs[0].text
    if "Carla Ferrante" not in riga or "Imperia" not in riga:
        raise AssertionError("i valori non sono nel documento: %r" % riga)


@prova("una cella che contiene un campo modulo non e' anche una cella vuota")
def _niente_doppio_posto():
    """Il campo mostra cinque spazi, quindi la cella sembra vuota: lo stesso
    posto risultava due volte, una come campo modulo e una come cella. Su un
    modulo del Demanio erano 162."""
    import docx
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    d = docx.Document()
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Ragione sociale"
    p = t.cell(0, 1).paragraphs[0]
    for xml in ('<w:r %s><w:fldChar w:fldCharType="begin"><w:ffData>'
                '<w:name w:val="T"/><w:enabled/><w:textInput/></w:ffData>'
                '</w:fldChar></w:r>',
                '<w:r %s><w:fldChar w:fldCharType="separate"/></w:r>',
                '<w:r %s><w:t xml:space="preserve">     </w:t></w:r>',
                '<w:r %s><w:fldChar w:fldCharType="end"/></w:r>'):
        p._p.append(parse_xml(xml % nsdecls("w")))
    fuori = io.BytesIO()
    d.save(fuori)
    percorso = TMP / "cella_modulo.docx"
    percorso.write_bytes(fuori.getvalue())
    m = word.leggi(percorso, TMP)
    tipi = [a["tipo"] for a in m["ancore"]]
    if tipi.count("cella") and tipi.count("modulo"):
        raise AssertionError("lo stesso posto conta due volte: %s" % tipi)
    if "modulo" not in tipi:
        raise AssertionError("dentro la cella il campo modulo non si vede: %s" % tipi)


@prova("una proposta fuori contratto non butta via tutto il modulo")
def _fuori_contratto():
    """Il guasto: il modello nomina un posto che non esiste, `abbina_pagina`
    solleva, e il modulo intero non produce niente. Successo su
    02_tricase_pitturazione_immobili: zero campi per una riga sbagliata, e
    quella riga non si poteva scrivere comunque perche' il posto non c'era."""
    from unittest.mock import patch
    from motore import agente

    risposta = ('```json\n[{"posto": 1, "campo": "nome"},'
                ' {"posto": 99, "campo": "nome"},'
                ' {"posto": 2, "campo": "inventata"}]\n```')
    indirizzi = {"1": "a1", "2": "a2"}

    with patch.object(agente, "_chiedi", return_value=risposta):
        try:
            agente.abbina_pagina("x", indirizzi, ["nome"])
        except agente.RispostaModelloInvalida:
            pass
        else:
            raise AssertionError("senza `scartate` deve restare severa")

    scartate = []
    with patch.object(agente, "_chiedi", return_value=risposta):
        deciso = agente.abbina_pagina("x", indirizzi, ["nome"], scartate=scartate)
    if deciso != {"a1": "nome"}:
        raise AssertionError("la proposta buona non e' sopravvissuta: %s" % deciso)
    if len(scartate) != 2:
        raise AssertionError("attese 2 proposte scartate, %d: %s"
                             % (len(scartate), scartate))


@prova("due chiavi diverse sullo stesso posto lasciano il posto vuoto")
def _due_chiavi_stesso_posto():
    """Non si sceglie a caso fra due: si lascia vuoto. Un campo vuoto lo
    riempie chi legge, un campo con la chiave sbagliata gli fa buttare via
    la bozza."""
    from unittest.mock import patch
    from motore import agente
    risposta = ('```json\n[{"posto": 1, "campo": "nome"},'
                ' {"posto": 1, "campo": "cognome"}]\n```')
    scartate = []
    with patch.object(agente, "_chiedi", return_value=risposta):
        deciso = agente.abbina_pagina("x", {"1": "a1"}, ["nome", "cognome"],
                                      scartate=scartate)
    if deciso:
        raise AssertionError("ha scelto una delle due invece di lasciare vuoto: %s"
                             % deciso)


@prova("se la rilettura non riesce, non si scrive - ma il modulo non salta")
def _rilettura_che_non_riesce():
    """Senza passare `errori`, un guasto nella rilettura - risposta
    illeggibile, rete che cade - fa saltare tutto il modulo con un
    traceback. Con `errori`, il meccanismo gia' scritto sospende le
    assegnazioni: non si scrive quello che non si e' potuto rileggere."""
    from unittest.mock import patch
    from motore import agente
    from motore.anagrafe import Anagrafe

    rese = [{"pagina": 1, "testo": "Nome «1:20»", "indirizzi": {"1": "a1"}}]
    anagrafe = Anagrafe({"profilo": {"nome_cognome": "Carla Ferrante"}})
    scelte = {"a1": "nome_cognome"}

    with patch.object(agente, "rileggi_pagina", side_effect=RuntimeError("caduta")):
        try:
            agente.rileggi_documento(rese, scelte, anagrafe)
        except RuntimeError:
            pass
        else:
            raise AssertionError("senza `errori` deve ancora sollevare")

        guasti = []
        bocciati = agente.rileggi_documento(rese, scelte, anagrafe, errori=guasti)
    if not guasti:
        raise AssertionError("il guasto non e' stato registrato")
    if bocciati.get("a1") is None:
        raise AssertionError("l'assegnazione non riletta doveva restare sospesa, "
                             "invece si sarebbe scritta: %s" % bocciati)

    # e la lavorazione lo passa davvero
    import inspect
    from motore import lavorazione_word
    sorgente = inspect.getsource(lavorazione_word.lavora)
    dove = sorgente.find("rileggi_documento")
    if "errori=" not in sorgente[dove:dove + 200]:
        raise AssertionError("la lavorazione Word non passa `errori`:\n%s"
                             % sorgente[dove:dove + 200])


@prova("una risposta illeggibile si richiede una volta, non due")
def _ritentativo():
    """Su 02_Demanio_Rimini la risposta e' arrivata illeggibile e il modulo
    intero e' andato perso; rifacendo la stessa identica domanda e' tornato
    un JSON pulito. E' un guasto di passaggio. Ma se sbaglia due volte
    uguale non lo e', e allora ci si ferma invece di spendere all'infinito."""
    import io as _io
    from unittest.mock import patch
    from motore import agente, fascicoli, lavorazione_word

    cartella = TMP / "ritentativo"
    percorso = cartella / "modulo.docx"
    percorso.parent.mkdir(parents=True, exist_ok=True)
    percorso.write_bytes(_documento(["Il sottoscritto ____________"]))
    fascicolo = {"profilo": {"nome_cognome": "Carla Ferrante"}}

    def corsa(risposte):
        passi = []
        with patch.object(agente, "_chiedi", side_effect=list(risposte)):
            for p in lavorazione_word.lavora(TMP, percorso, fascicolo,
                                             cartella, pertinenza_confermata=True,
                                             condizioni=set()):
                passi.append(p)
        return passi

    buona = '```json\n[{"posto": 1, "campo": "nome_cognome"}]\n```'
    # prima illeggibile, poi buona: si arriva in fondo. La terza risposta e'
    # per la rilettura, che a questo punto parte.
    passi = corsa(["non sono json", buona, "[]"])
    if passi[-1].get("verdetto") is None:
        raise AssertionError("dopo il secondo tentativo doveva compilare: %s"
                             % passi[-1]["testo"][:90])

    # due volte illeggibile: si ferma, e lo dice
    passi = corsa(["non sono json", "nemmeno io"])
    if passi[-1].get("verdetto") is not None:
        raise AssertionError("doveva fermarsi dopo due risposte illeggibili")
    if "illeggibile" not in passi[-1]["testo"] and "leggere" not in passi[-1]["testo"]:
        raise AssertionError("non spiega perche' si e' fermato: %r"
                             % passi[-1]["testo"][:90])


@prova("i trattini troppo corti non sono campi")
def _troppo_corti():
    m = _leggi(_documento(["parola_con_trattino e __ due soli"]), TMP)
    if [a for a in m["ancore"] if a["tipo"] == "riempimento"]:
        raise AssertionError("ha preso per campo un trattino corto")


if __name__ == "__main__":
    falliti = 0
    for descrizione, funzione in _prove:
        try:
            funzione()
            print("  ok      %s" % descrizione)
        except AssertionError as guaio:
            falliti += 1
            print("  FALLITA %s\n            %s" % (descrizione, guaio))
    print("\n%d prove, %d fallite" % (len(_prove), falliti))
    raise SystemExit(1 if falliti else 0)
