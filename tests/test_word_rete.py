"""Un campo che non regge non deve far perdere il modulo, e i controlli Word
posti su un paragrafo o una cella propri si scrivono come quelli in linea."""
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from motore import word, word_inventory, word_patch, agente
from word_fixtures import FIELDS, paragraphs

SEGNAPOSTO = 'Fare clic qui per immettere testo.'


def _controllo_blocco(alias, rpr_controllo=True):
    rpr = ('<w:rPr><w:rFonts w:ascii="Arial Narrow" w:hAnsi="Arial Narrow"/>'
           '<w:sz w:val="24"/></w:rPr>') if rpr_controllo else ''
    return parse_xml(
        '<w:sdt %s><w:sdtPr>%s<w:alias w:val="%s"/><w:lock w:val="sdtLocked"/>'
        '<w:showingPlcHdr/></w:sdtPr><w:sdtContent><w:p><w:r><w:rPr>'
        '<w:rStyle w:val="PlaceholderText"/></w:rPr><w:t>%s</w:t></w:r></w:p>'
        '</w:sdtContent></w:sdt>' % (nsdecls('w'), rpr, alias, SEGNAPOSTO))


def blocchi(d):
    """Controlli su un paragrafo proprio: prima facevano fallire tutto."""
    for label, _ in FIELDS:
        d.add_paragraph(label)
        d.element.body.insert(len(d.element.body) - 1, _controllo_blocco(label))


def celle(d):
    """Controlli che occupano una cella intera."""
    t = d.add_table(rows=len(FIELDS), cols=2)
    for row, (label, _) in zip(t.rows, FIELDS):
        row.cells[0].text = label
        tc = row.cells[1]._tc
        for p in tc.findall(qn('w:p')):
            tc.remove(p)
        tc.append(_controllo_blocco(label))


class ReteTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.network = patch.object(agente, '_chiedi', side_effect=AssertionError('No network in tests'))
        self.network.start()

    def tearDown(self):
        self.network.stop()
        self.temp.cleanup()

    def source(self, build, name='source'):
        doc = Document(); build(doc)
        path = self.root / (name + '.docx'); doc.save(path)
        return path

    def fill(self, source, values_by_index=None):
        m = word.leggi(source)
        ids = [a['id'] for a in m['ancore'] if a['scrivibile']]
        values = {i: FIELDS[n % len(FIELDS)][1] for n, i in enumerate(ids)}
        out = self.root / 'filled.docx'
        rep = word.scrivi(source, m, values, out)
        return m, out, rep, word.verifica(out, rep)

    def test_controlli_su_paragrafo_o_cella_si_scrivono(self):
        for nome, build in (('blocchi', blocchi), ('celle', celle)):
            with self.subTest(nome):
                m, out, rep, check = self.fill(self.source(build, nome))
                self.assertEqual(sum(a['scrivibile'] for a in m['ancore']), 4)
                self.assertEqual(len(check['verificate']), 4)
                self.assertFalse(rep['mancate'])
                _, roots = word_inventory.apri(out.read_bytes())
                for sdt in roots['word/document.xml'].iter(qn('w:sdt')):
                    pr = sdt.find(qn('w:sdtPr'))
                    self.assertIsNone(pr.find(qn('w:showingPlcHdr')))
                    self.assertEqual(pr.find(qn('w:lock')).get(qn('w:val')), 'sdtLocked')

    def test_il_valore_prende_lo_stile_del_controllo_non_del_segnaposto(self):
        m, out, rep, check = self.fill(self.source(blocchi))
        _, roots = word_inventory.apri(out.read_bytes())
        for sdt in roots['word/document.xml'].iter(qn('w:sdt')):
            run = sdt.find('.//' + qn('w:r'))
            self.assertIsNone(run.find('.//' + qn('w:rStyle')))
            font = run.find('w:rPr/w:rFonts', word_inventory.NS)
            self.assertEqual(font.get(qn('w:ascii')), 'Arial Narrow')

    def test_senza_stile_del_controllo_si_toglie_solo_il_segnaposto(self):
        def build(d):
            d.add_paragraph('Nome')
            sdt = _controllo_blocco('Nome', rpr_controllo=False)
            sdt.find('.//' + qn('w:rPr')).append(OxmlElement('w:b'))
            d.element.body.insert(len(d.element.body) - 1, sdt)
        m, out, rep, check = self.fill(self.source(build))
        self.assertEqual(len(check['verificate']), 1)
        _, roots = word_inventory.apri(out.read_bytes())
        run = next(roots['word/document.xml'].iter(qn('w:sdt'))).find('.//' + qn('w:r'))
        self.assertIsNone(run.find('.//' + qn('w:rStyle')))
        self.assertIsNotNone(run.find('w:rPr/w:b', word_inventory.NS))

    def test_senza_guasti_il_ripiego_e_identico_alla_scrittura(self):
        src = self.source(lambda d: paragraphs(d, '_' * 25))
        data = src.read_bytes()
        fisico = word_inventory.estrai(data)
        valori = {p['riferimento']: 'Mario Rossi' for p in fisico['punti'] if p['scrivibile']}
        diretto, _ = word_patch.scrivi(data, fisico, valori)
        ripiego, _, scartati = word_patch.scrivi_con_ripiego(data, fisico, valori)
        self.assertEqual(diretto, ripiego)
        self.assertEqual(scartati, {})

    def _con_guasto(self, marcatore):
        """Una scrittura che rompe la struttura quando il paragrafo contiene il marcatore."""
        originale = word_patch._intervallo

        def rotto(p, inizio, fine, valore):
            originale(p, inizio, fine, valore)
            if marcatore in word_inventory.testo_ancore(p)[0]:
                p.addnext(OxmlElement('w:p'))
        return patch.object(word_patch, '_intervallo', side_effect=rotto)

    def test_un_campo_che_rompe_non_fa_perdere_il_modulo(self):
        src = self.source(lambda d: paragraphs(d, '_' * 25))
        m = word.leggi(src)
        ids = [a['id'] for a in m['ancore']]
        valori = dict(zip(ids, ['Uno', 'ROMPE', 'Tre', 'Quattro']))
        out = self.root / 'filled.docx'
        with self._con_guasto('ROMPE'):
            rep = word.scrivi(src, m, valori, out)
        check = word.verifica(out, rep)
        self.assertEqual(sorted(w['valore'] for w in check['verificate']), ['Quattro', 'Tre', 'Uno'])
        self.assertEqual([x['ancora'] for x in rep['mancate']], [ids[1]])
        self.assertIn('da compilare a mano', rep['mancate'][0]['motivo'])

    def test_la_prova_in_lettura_toglie_il_campo_prima_del_modello(self):
        def build(d):
            d.add_paragraph('Ragione sociale: ' + '_' * 25)
            d.add_paragraph('GUASTO: ' + '_' * 25)
            d.add_paragraph('PEC: ' + '_' * 25)
        src = self.source(build)
        with self._con_guasto('GUASTO'):
            m = word.leggi(src)
        stato = {a['etichetta']: a['scrivibile'] for a in m['ancore']}
        self.assertEqual(stato, {'Ragione sociale': True, 'GUASTO': False, 'PEC': True})
        motivo = next(a['motivo'] for a in m['ancore'] if a['etichetta'] == 'GUASTO')
        self.assertIn('da compilare a mano', motivo)
        testo, indirizzi = word.rendi(m)
        self.assertEqual(len(indirizzi), 2)


if __name__ == '__main__':
    unittest.main()
