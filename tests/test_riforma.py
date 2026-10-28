"""General structure and workflow invariants, synthetic data, no network."""
import copy
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import fitz
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from requests import Timeout
from streamlit.testing.v1 import AppTest
from motore import agente, classificatore, geometria, controllo, compilatore, word
from motore import lavori, lavorazione, piano_pdf, fascicoli
import esecuzione_sorvegliata as legacy

ROOT = Path(__file__).resolve().parents[1]


class Riforma(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.network = patch.object(agente, '_chiedi', side_effect=AssertionError('No network'))
        self.network.start()

    def tearDown(self):
        self.network.stop()
        self.tmp.cleanup()

    def source(self):
        p = self.root/'source.pdf'
        with fitz.open() as d:
            page = d.new_page()
            for i, label in enumerate(('Campo uno', 'Campo due')):
                y=60+i*40
                page.insert_text((30,y),label)
                page.draw_rect((140,y-14,340,y+10))
            d.save(p)
        return p

    def test_grid_edges_across_translations_scales_and_open_sides(self):
        for scale in (0.5,1,2):
            for dx,dy in ((0,0),(47,103)):
                for opened in (('right',),('left','right','top'),('top','bottom'),('left','right','top','bottom')):
                    with self.subTest(scale=scale, translation=(dx,dy), opened=opened):
                        xs=[dx+scale*x for x in (10,70,130,190,250)]
                        ys=[dy+scale*y for y in (20,40,60,80,100)]
                        v=[geometria.Segmento(x,ys[0],ys[-1]) for i,x in enumerate(xs)
                           if not (i==0 and 'left' in opened) and not(i==len(xs)-1 and 'right' in opened)]
                        h=[geometria.Segmento(y,xs[0],xs[-1]) for i,y in enumerate(ys)
                           if not(i==0 and 'top' in opened) and not(i==len(ys)-1 and 'bottom' in opened)]
                        vv,hh,_=geometria.completa_griglie(v,h)
                        self.assertEqual(len(geometria._celle(vv,hh)),16)

    def test_page_separators_are_not_invented_into_a_table(self):
        h=[geometria.Segmento(y,0,500) for y in (10,30,50,70)]
        v=[geometria.Segmento(50,10,70)]
        vv,hh,inferred=geometria.completa_griglie(v,h)
        self.assertFalse(inferred)
        self.assertFalse(geometria._celle(vv,hh))

    def test_neighbouring_tables_are_not_joined(self):
        v=[];h=[]
        for offset in (0,250):
            v += [geometria.Segmento(offset+x,20,100) for x in (30,60,90)]
            h += [geometria.Segmento(y,offset,offset+120) for y in (20,40,60,80,100)]
        vv,hh,_=geometria.completa_griglie(v,h)
        self.assertTrue(all(c.riquadro.width<=30 for c in geometria._celle(vv,hh)))

    def test_same_row_label_beats_previous_heading(self):
        with fitz.open() as d:
            p=d.new_page();p.insert_text((40,40),'Titolo lontano')
            p.insert_text((40,85),'Etichetta corretta')
            p.draw_rect((30,65,140,95));p.draw_rect((140,65,300,95))
            g=geometria.leggi(p,0)
            a=next(a for a in classificatore.classifica_pagina(g) if a.tipo=='area')
            self.assertEqual(a.etichetta,'Etichetta corretta')

    def test_empty_pdf_is_not_complete(self):
        source=self.source();m=piano_pdf.mappe.costruisci(source)
        v=controllo.controlla(source,source,m,{'scritture':[]},{})
        self.assertNotEqual(v['stato'],controllo.COMPLETO)
        self.assertEqual(len(v['punti_aperti']),len(m['ancore']))
        self.assertIn('non risulta compilato',controllo.racconta(v))

    def test_render_cache_is_bound_to_document_lifetime(self):
        from riusare.strumenti_documento import _chiave, inchiostro
        identities=[]
        for i in range(20):
            with fitz.open() as d:
                page=d.new_page()
                if i%2:page.draw_rect((50,50,100,100),fill=(0,0,0))
                identities.append(_chiave(page,150))
                for _ in range(5):
                    _,pixels=inchiostro(page,(55,55,95,95),dpi=150)
                    self.assertEqual(pixels>0,bool(i%2))
        self.assertEqual(len(set(identities)),len(identities))

    def test_legacy_exclusion_preserves_other_field(self):
        code='def valori(fascicolo):\n    return {"a":"Alfa", "b":"Beta"}\n'+compilatore.SEGNAPOSTO
        values,_=legacy.scegli_valori(compilatore.senza(code,['a']),{}, {'a':{},'b':{}})
        self.assertEqual(values,{'b':'Beta'})

    def test_legacy_rejects_object_attributes_and_keeps_data_access(self):
        with self.assertRaises(legacy.ValoriInvalidi):
            legacy.scegli_valori('def valori(fascicolo):\n    return {"a": fascicolo.__class__}',{}, {'a':{}})
        v,_=legacy.scegli_valori('def valori(fascicolo):\n    p = fascicolo["profilo"]\n    return {"a":p.get("nome").strip()}',{'profilo':{'nome':' Alfa '}},{'a':{}})
        self.assertEqual(v,{'a':'Alfa'})

    def test_conflicting_assignments_cannot_reenter(self):
        response=json.dumps([{'posto':1,'campo':k} for k in ('nome','pec','nome')])
        with patch.object(agente,'_chiedi',return_value=response):
            self.assertEqual(agente.abbina_pagina('',{'1':'a'},['nome','pec'],scartate=[]),{})
        with patch.object(agente,'_chiedi',return_value='[{"posto":1,"campo":"nome"},{"posto":1,"campo":""}]'):
            self.assertEqual(agente.abbina_pagina('',{'1':'a'},['nome'],scartate=[],decisioni={}),{})

    def test_declarative_contract_rejects_duplicates_unknown_keys_and_omissions(self):
        response=json.dumps([dict(posto=1,campo='nome',stato='assegna',soggetto='Impresa',motivo='Etichetta'),
                             dict(posto=1,campo='',stato='manuale'),
                             dict(posto=2,campo='inventata',stato='assegna')])
        with patch.object(agente,'_chiedi',return_value=response):
            d=agente.decidi_campi('',{'1':'a','2':'b','3':'c'},[],['nome'])
        self.assertEqual(set(d),{'a','b','c'})
        self.assertTrue(all(v['stato']=='dubbio' for v in d.values()))

    def test_verified_plan_reuses_bindings_only_for_unchanged_profile(self):
        source=self.source();profile={'profilo':{'nome':'Alfa'}}
        def decide(text,addresses,*args,**kwargs):
            return {a:dict(stato='assegna',campo='nome',soggetto='Impresa',motivo='Etichetta') for a in addresses.values()}
        with patch.object(agente,'disponibile',return_value=True),patch.object(agente,'immagini',return_value=[]),patch.object(agente,'decidi_campi',side_effect=decide) as call,patch.object(agente,'rileggi_documento',return_value={}):
            a=list(lavorazione.lavora(self.root,source,profile,self.root/'runs',pertinenza_confermata=True))[-1]
            b=list(lavorazione.lavora(self.root,source,profile,self.root/'runs'))[-1]
            self.assertEqual(call.call_count,1)
            self.assertNotEqual(a['bozza'],b['bozza'])
            changed={'profilo':{'nome':'Beta'}}
            c=list(lavorazione.lavora(self.root,source,changed,self.root/'runs',pertinenza_confermata=True))[-1]
            self.assertEqual(call.call_count,2)
            self.assertEqual({r['valore'] for r in c['verdetto']['verificate']},{'Beta'})

    def test_pertinence_unknown_requires_confirmation(self):
        for response in ('incomplete','[]','null','{}'):
            with patch.object(agente,'_chiedi',return_value=response):
                self.assertFalse(agente.pertinenza('',{'profilo':{}})['pertinente'])

    def test_failed_conversion_never_returns_old_file(self):
        from subprocess import CompletedProcess
        src=self.root/'m.doc';src.write_bytes(b'old format')
        out=self.root/'conv';out.mkdir();(out/'m.docx').write_bytes(b'previous')
        with patch.object(word,'_soffice',return_value='mock'),patch.object(word.subprocess,'run',return_value=CompletedProcess([],1)):
            with self.assertRaises(word.ConversioneAssente):word.converti(src,out)

    def test_preexisting_legacy_formtext_is_preserved(self):
        doc=Document();p=doc.add_paragraph('Nome: ')
        for tag,value in [('fldChar','begin'),('instrText',' FORMTEXT '),('fldChar','separate'),('t','Gia compilato'),('fldChar','end')]:
            el=OxmlElement('w:'+tag)
            if tag=='fldChar':el.set(qn('w:fldCharType'),value)
            else:el.text=value
            p.add_run()._r.append(el)
        src=self.root/'filled.docx';doc.save(src)
        m=word.leggi(src)
        self.assertTrue(m['ancore'])
        self.assertFalse(any(a['scrivibile'] for a in m['ancore']))
        rendered,addresses=word.rendi(m)
        self.assertFalse(addresses)
        self.assertIn('Gia compilato',rendered)
        for el in doc.element.iter(qn('w:t')):
            if el.text == 'Gia compilato':el.text='__/__/____'
        blank=self.root/'date_blank.docx';doc.save(blank)
        self.assertTrue(any(a['scrivibile'] for a in word.leggi(blank)['ancore']))

    def test_sdt_wrapped_row_and_cell_reading(self):
        for tag in ('tr','tc'):
            doc=Document();t=doc.add_table(rows=1,cols=2);t.cell(0,0).text='Nome';t.cell(0,1).text='________'
            node=t.rows[0]._tr if tag=='tr' else t.cell(0,1)._tc
            parent=node.getparent();idx=parent.index(node)
            wrapper=OxmlElement('w:sdt');wrapper.append(OxmlElement('w:sdtPr'))
            content=OxmlElement('w:sdtContent');wrapper.append(content)
            parent.remove(node);content.append(node);parent.insert(idx,wrapper)
            src=self.root/(tag+'.docx');doc.save(src)
            self.assertTrue(word.leggi(src)['ancore'])

    def test_profile_switch_clears_widget_and_results(self):
        state={}
        lavori.sincronizza_fascicolo(state,'a.json',{'profilo':{'nome':'A'}})
        state.update(ana_nome='Unsaved A',w_verdetto={'x':1})
        lavori.sincronizza_fascicolo(state,'b.json',{'profilo':{'nome':'B'}})
        self.assertNotIn('ana_nome',state)
        self.assertIsNone(state['w_verdetto'])

    def test_declarative_plan_correction_exclusion_and_source_binding(self):
        source=self.source();m=piano_pdf.mappe.costruisci(source)
        ids=[a['id'] for a in m['ancore']]
        profile={'profilo':{'nome':'Alfa','pec':'info@example.invalid'}}
        plan=dict(versione=piano_pdf.VERSIONE,impronta=m['impronta'],fascicolo=lavori.impronta_dati(profile),
                  mappa=piano_pdf._firma_mappa(m),decisioni={a:dict(stato='assegna',campo='nome',soggetto='Impresa',motivo='Test') for a in ids})
        a=piano_pdf.esegui(source,m,plan,profile,lavori.nuova_cartella(self.root/'runs'))
        b=piano_pdf.esegui(source,m,a['verdetto']['piano'],profile,lavori.nuova_cartella(self.root/'runs'),togli=[ids[0]])
        self.assertNotEqual(a['bozza'],b['bozza'])
        self.assertEqual(len(a['verdetto']['verificate']),2)
        self.assertEqual([x['ancora'] for x in b['verdetto']['verificate']],[ids[1]])
        c=piano_pdf.esegui(source,m,b['verdetto']['piano'],profile,lavori.nuova_cartella(self.root/'runs'),togli=[ids[1]])
        self.assertFalse(c['verdetto']['verificate'])
        changed=copy.deepcopy(profile);changed['profilo']['nome']='Beta'
        with self.assertRaises(ValueError):piano_pdf.esegui(source,m,plan,changed,self.root)

    def test_failed_decision_is_not_cached_and_other_pages_are_diagnosed(self):
        source=self.source();profile={'profilo':{'nome':'Alfa'}}
        with patch.object(agente,'disponibile',return_value=True),patch.object(agente,'immagini',return_value=[]),patch.object(agente,'decidi_campi',side_effect=Timeout()):
            events=list(lavorazione.lavora(self.root,source,profile,self.root/'runs',pertinenza_confermata=True))
        v=events[-1]
        self.assertEqual(v['passo'],'rifiuto')
        self.assertNotIn('bozza',v)
        self.assertEqual(v['guasti_lettura'],[1])
        self.assertFalse(list((self.root/'runs'/'piani').glob('*.json')))

    def test_user_exclusion_and_correction_persist_across_compilations(self):
        source=self.source();m=piano_pdf.mappe.costruisci(source)
        profile={'profilo':{'nome':'Alfa','altro':'Beta'}}
        plan=dict(versione=piano_pdf.VERSIONE,impronta=m['impronta'],fascicolo=lavori.impronta_dati(profile),
                  mappa=piano_pdf._firma_mappa(m),decisioni={a['id']:dict(stato='assegna',campo='nome',soggetto='Impresa',motivo='Test') for a in m['ancore']})
        aid=m['ancore'][0]['id']; other=m['ancore'][1]['id']
        plan['decisioni'][other]['campo']='altro'
        first=list(lavorazione.lavora(self.root,source,profile,self.root/'runs',piano_di_prima=plan,togli=[aid]))[-1]
        second=list(lavorazione.lavora(self.root,source,profile,self.root/'runs'))[-1]
        self.assertEqual(second['verdetto']['piano']['decisioni'],first['verdetto']['piano']['decisioni'])
        self.assertEqual([r['valore'] for r in second['verdetto']['verificate']],['Beta'])

    def test_user_override_does_not_cache_failed_page(self):
        source=self.source()
        with fitz.open(source) as doc:
            doc.fullcopy_page(0)
            two=self.root/'two.pdf';doc.save(two)
        profile={'profilo':{'nome':'Alfa'}}
        calls=[]
        def decide(text, addresses, *args, **kwargs):
            calls.append(1)
            if 'p2.' in next(iter(addresses.values())):raise Timeout()
            return {a:dict(stato='assegna',campo='nome',soggetto='Impresa',motivo='Test') for a in addresses.values()}
        with patch.object(agente,'disponibile',return_value=True),patch.object(agente,'immagini',return_value=[]),patch.object(agente,'decidi_campi',side_effect=decide),patch.object(agente,'rileggi_documento',return_value={}):
            first=list(lavorazione.lavora(self.root,two,profile,self.root/'runs',pertinenza_confermata=True))[-1]
            plan=first['verdetto']['piano'];aid=next(a for a in plan['decisioni'] if a.startswith('p1.'))
            second=list(lavorazione.lavora(self.root,two,profile,self.root/'runs',piano_di_prima=plan,togli=[aid]))[-1]
            cache=list((self.root/'runs'/'piani').glob('*.json'))
            self.assertTrue(cache)
            self.assertTrue(all(p.name.endswith('.utente.json') for p in cache))
            count=len(calls)
            third=list(lavorazione.lavora(self.root,two,profile,self.root/'runs',pertinenza_confermata=True))[-1]
            self.assertGreater(len(calls),count)
            self.assertNotIn(aid,[r['ancora'] for r in third['verdetto']['verificate']])

    def test_failed_user_correction_cannot_restore_previous_value(self):
        source=self.source();m=piano_pdf.mappe.costruisci(source)
        profile={'profilo':{'nome':'Alfa','lungo':'W'*5000}}
        plan=dict(versione=piano_pdf.VERSIONE,impronta=m['impronta'],fascicolo=lavori.impronta_dati(profile),
                  mappa=piano_pdf._firma_mappa(m),decisioni={a['id']:dict(stato='assegna',campo='nome',soggetto='Impresa',motivo='Utente',origine='utente') for a in m['ancore']})
        aid=m['ancore'][0]['id']
        first=list(lavorazione.lavora(self.root,source,profile,self.root/'runs',piano_di_prima=plan))[-1]
        changed=copy.deepcopy(first['verdetto']['piano']);changed['decisioni'][aid]['campo']='lungo'
        second=list(lavorazione.lavora(self.root,source,profile,self.root/'runs',piano_di_prima=changed))[-1]
        self.assertIn(aid,second['verdetto']['scritture_escluse'])
        saved=json.loads(next((self.root/'runs'/'piani').glob('*.utente.json')).read_text(encoding='utf-8'))
        self.assertEqual(saved['decisioni'][aid]['stato'],'manuale')
        self.assertEqual(saved['decisioni'][aid]['campo'],'lungo')

    def test_app_profile_switch_does_not_copy_unsaved_fields(self):
        folder=self.root/'fascicoli';folder.mkdir()
        for name in ('a','b'):fascicoli.salva(folder/(name+'.json'),{'profilo':{'ragione_sociale':name}})
        with patch.dict(os.environ,{'ACCESSO_LIBERO':'1','APP_SCADENZA':'NESSUNA','CLIENTI':'','QUADRA_DATI':str(self.root)}):
            app=AppTest.from_file(str(ROOT/'app.py'),default_timeout=30).run()
            app.text_input(key='ana_ragione_sociale').input('Unsaved A').run()
            app.selectbox[0].select(folder/'b.json').run()
            self.assertEqual(app.text_input(key='ana_ragione_sociale').value,'b')
            self.assertFalse(app.exception)


if __name__=='__main__':unittest.main()
